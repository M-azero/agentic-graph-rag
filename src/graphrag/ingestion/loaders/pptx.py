"""PowerPoint (.pptx) loader: slide text, tables and speaker notes.

One block per slide, prefixed with the slide number. Deck text is short and
repeats heavily between slides — the same title, the same footer — so without
that marker every retrieved chunk looks alike and a citation cannot say where
it came from.

Speaker notes are included because they are usually where the sentence explaining
the slide actually lives; a bullet list of three words is not much to retrieve on.
"""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from typing import Any

from graphrag.core.errors import IngestionError
from graphrag.core.types import Document
from graphrag.ingestion.loaders.base import Loader


def _shape_text(shape: Any) -> Iterator[str]:
    """Text from one shape, recursing into groups.

    Grouped shapes are ordinary in a deck and hold their own `.shapes`, so a
    flat pass over the top level silently drops whole diagrams.
    """
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            line = " | ".join(cell.text.strip() for cell in row.cells)
            if line.strip(" |"):
                yield line
        return
    nested = getattr(shape, "shapes", None)
    if nested is not None:
        for inner in nested:
            yield from _shape_text(inner)
        return
    if getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text.strip()
        if text:
            yield text


class PPTXLoader(Loader):
    suffixes = (".pptx",)

    def load(self, path: Path) -> Document:
        try:
            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover
            raise IngestionError("python-pptx is not installed") from exc

        deck = Presentation(str(path))
        blocks: list[str] = []
        for number, slide in enumerate(deck.slides, start=1):
            parts = [text for shape in slide.shapes for text in _shape_text(shape)]
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"Notes: {notes}")
            if parts:
                blocks.append(f"[Slide {number}]\n" + "\n".join(parts))
        return Document(
            source=str(path),
            content="\n\n".join(blocks),
            metadata={"type": "pptx", "slides": len(deck.slides)},
        )
