"""Excel (.xlsx) loader. Rows become labeled sentence blocks, one per row.

The same shaping as `CSVLoader`, for the same reason — an unlabeled grid of
values embeds poorly — with the two differences a workbook forces: every block
carries its sheet name, so a retrieved chunk says which table it came from, and
each sheet gets its own header row.

Only the OOXML formats are claimed. `.xls` is a different, binary format that
openpyxl cannot open, so listing it here would trade a clear "unsupported file"
at upload for a parser error a minute into the ingest job.
"""

from __future__ import annotations

from pathlib import Path

from graphrag.core.errors import IngestionError
from graphrag.core.types import Document
from graphrag.ingestion.loaders.base import Loader

_MAX_ROWS = 5000  # across the whole workbook, matching CSVLoader


class XLSXLoader(Loader):
    suffixes = (".xlsx", ".xlsm")

    def load(self, path: Path) -> Document:
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover
            raise IngestionError("openpyxl is not installed") from exc

        # read_only streams rows instead of materializing every sheet, and
        # data_only takes a formula's cached result — "=SUM(B2:B9)" is not what
        # anyone is searching for. The trade is that a workbook Excel has never
        # saved has no cached values, so those cells read empty.
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = len(book.sheetnames)
        blocks: list[str] = []
        try:
            for sheet in book.worksheets:
                rows = sheet.iter_rows(values_only=True)
                first = next(rows, None) or ()
                header = [str(c).strip() if c is not None else "" for c in first]
                for row in rows:
                    if len(blocks) >= _MAX_ROWS:
                        break
                    pairs = [
                        f"{(header[j] if j < len(header) and header[j] else f'col{j + 1}')}"
                        f": {str(cell).strip()}"
                        for j, cell in enumerate(row)
                        if cell is not None and str(cell).strip()
                    ]
                    if pairs:
                        blocks.append(f"[{sheet.title}] " + "; ".join(pairs))
        finally:
            book.close()  # read_only holds the zip handle open until told otherwise
        return Document(
            source=str(path),
            content="\n\n".join(blocks),
            metadata={"type": "xlsx", "sheets": sheets, "rows": len(blocks)},
        )
